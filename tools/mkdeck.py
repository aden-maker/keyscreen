"""Build the 答辩演示 .pptx (16:9, ~26 slides). Usage: mkdeck.py OUT.pptx [IMGDIR]"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "evidence" / "deck"
NAVY, BLUE, GREY, LIGHT, WHITE, RED = "1F3864", "2E75B6", "595959", "EEF3FA", "FFFFFF", "C00000"
W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BL = prs.slide_layouts[6]
n = [0]


def rgb(s):
    return RGBColor.from_string(s)


def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, x, y, w, h)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(sl, x, y, w, h, lines, size=18, color="222222", bold=False, align=PP_ALIGN.LEFT, bullet=False, gap=6):
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        b, sz, col = bold, size, color
        if isinstance(ln, tuple):
            d = (size, color, bold)
            ln, sz, col, b = ln + d[len(ln) - 1:]
        r = p.add_run()
        r.text = ("• " if bullet and ln else "") + ln
        r.font.size, r.font.bold, r.font.name = Pt(sz), b, "微软雅黑"
        r.font.color.rgb = rgb(col)
    return tb


def slide(title, sub=None):
    sl = prs.slides.add_slide(BL)
    n[0] += 1
    box(sl, 0, 0, W, Inches(1.05), NAVY)
    box(sl, 0, Inches(1.05), W, Inches(0.06), BLUE)
    text(sl, Inches(0.5), Inches(0.18), Inches(11), Inches(0.8), title, 28, WHITE, True)
    if sub:
        text(sl, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.5), sub, 14, GREY)
    text(sl, Inches(11.3), Inches(7.05), Inches(1.8), Inches(0.35), f"无敌键队 · {n[0]}", 10, GREY, align=PP_ALIGN.RIGHT)
    return sl


def bullets(sl, items, x=0.6, y=1.75, w=12.1, h=5.2, size=18):
    text(sl, Inches(x), Inches(y), Inches(w), Inches(h), items, size, bullet=True, gap=10)


def cards(sl, items, y=1.9, h=4.6, cols=3, size=14):
    gw, gap = (12.3 - 0.3 * (cols - 1)) / cols, 0.3
    for i, (t, body) in enumerate(items):
        x = 0.5 + (i % cols) * (gw + gap)
        yy = y + (i // cols) * (h + 0.25)
        box(sl, Inches(x), Inches(yy), Inches(gw), Inches(h), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        box(sl, Inches(x), Inches(yy), Inches(0.12), Inches(h), BLUE)
        text(sl, Inches(x + 0.25), Inches(yy + 0.15), Inches(gw - 0.4), Inches(0.6), t, 17, NAVY, True)
        text(sl, Inches(x + 0.25), Inches(yy + 0.75), Inches(gw - 0.4), Inches(h - 0.9), body, size, bullet=True)


def stats(sl, items, y=1.8):
    gw = 12.3 / len(items)
    for i, (v, lab) in enumerate(items):
        x = 0.5 + i * gw
        box(sl, Inches(x + 0.1), Inches(y), Inches(gw - 0.2), Inches(1.5), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, Inches(x + 0.1), Inches(y + 0.1), Inches(gw - 0.2), Inches(0.8), v, 30, WHITE, True, PP_ALIGN.CENTER)
        text(sl, Inches(x + 0.1), Inches(y + 0.9), Inches(gw - 0.2), Inches(0.5), lab, 13, "D9E2F3", align=PP_ALIGN.CENTER)


def table(sl, rows, x=0.5, y=1.8, w=12.3, size=13, widths=None):
    t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows))).table
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            cell.text = c
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name = Pt(size), "微软雅黑"
                    run.font.bold = i == 0
                    run.font.color.rgb = rgb(WHITE if i == 0 else "222222")
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else (LIGHT if i % 2 else WHITE))
    if widths:
        for j, wd in enumerate(widths):
            t.columns[j].width = Inches(wd)


def pic(sl, name, x, y, w=None, h=None, cap=None):
    p = IMG / name
    if not p.exists():
        box(sl, Inches(x), Inches(y), Inches(w or 4), Inches(h or 3), LIGHT, GREY)
        text(sl, Inches(x), Inches(y + 1), Inches(w or 4), Inches(1), f"[缺图 {name}]", 12, GREY, align=PP_ALIGN.CENTER)
        return
    kw = {"width": Inches(w)} if w else {"height": Inches(h)}
    s = sl.shapes.add_picture(str(p), Inches(x), Inches(y), **kw)
    if cap:
        text(sl, Inches(x), Emu(s.top + s.height) + Inches(0.05), Emu(s.width), Inches(0.4), cap, 12, GREY, align=PP_ALIGN.CENTER)


def arrow_flow(sl, steps, y=2.0, h=1.1, size=13):
    gw = (12.3 - 0.25 * (len(steps) - 1)) / len(steps)
    for i, s in enumerate(steps):
        x = 0.5 + i * (gw + 0.25)
        box(sl, Inches(x), Inches(y), Inches(gw), Inches(h), BLUE if i % 2 == 0 else NAVY, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        text(sl, Inches(x + 0.25), Inches(y + 0.1), Inches(gw - 0.5), Inches(h - 0.2), s, size, WHITE, True, PP_ALIGN.CENTER)


# ---------------------------------------------------------------- slides
sl = prs.slides.add_slide(BL)
n[0] += 1
box(sl, 0, 0, W, H, NAVY)
box(sl, 0, Inches(4.9), W, Inches(0.08), BLUE)
text(sl, Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.2), "KeyScreen", 60, WHITE, True)
text(sl, Inches(0.8), Inches(2.7), Inches(11.5), Inches(1), "会被 AI 改造的多语言触屏键盘", 34, "D9E2F3", True)
text(sl, Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.8), "LubanCat 1H 触摸屏 × USB HID 键盘 × Agnes Harness（agnes-3.0-flash）", 18, "BDD7EE")
text(sl, Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.4),
     ["2026 江苏省 AI+科学与工程创新实践黑客松", "本科生组 · 参赛编号 U376 · 无敌键队 · 队长 贾灵鑫"], 18, WHITE, gap=8)

sl = slide("目录")
cards(sl, [("01 问题与目标", ["问题来源：三个真实痛点", "用户场景与目标", "方案总览"]),
           ("02 系统设计", ["系统架构", "USB HID 键盘实现", "板端软件与输入法"]),
           ("03 五大创新点", ["AI 动手改硬件", "一句话多步定制（AGH）", "四语言键盘 · 整句输入法 · 个性化校准"]),
           ("04 科学实验", ["基线数据", "留出法回放与诚实结论"]),
           ("05 工程与验证", ["可靠性设计", "实测指标", "11 条测试样例"]),
           ("06 价值与展望", ["解决了哪些问题", "实际应用场景", "社会价值 · 下一步 · 团队"])], h=2.3, size=13)

sl = slide("问题来源：三个真实痛点", "来自队员自己写外语作业、做项目演示时的切身体会")
cards(sl, [("① 外语键盘：买、贴、切，都麻烦", ["俄语/法语/德语键盘要额外购买或贴键帽贴纸", "系统切布局后键帽与实际输出不一致，新手频繁打错", "同一台电脑多人、多语言共用时更混乱"]),
           ("② 触屏键盘千人一面", ["每个人手指大小、落点习惯不同", "触屏键盘给所有人同一套固定键区", "板上实测：平均 22 px 系统性偏差，41% 按键落在键边缘附近"]),
           ("③ AI 只停留在屏幕里", ["大模型会聊天、会写字，但很少真正操作身边硬件", "缺少可审计的“看状态→动手→复核”闭环", "外设定制仍然依赖人手改配置文件"])])

sl = slide("目标与用户场景")
bullets(sl, [("目标", 20, NAVY, True),
             "一块触摸屏 = 一把即插即用的 USB 键盘，任何电脑免驱可用",
             "一块屏幕 = 四种语言键盘，键帽所见即所得，电脑布局自动跟随",
             "让 AI 能“看见”真实打字数据并动手改硬件：键位、语言、外观、桌宠",
             "一句话完成多步场景配置，每一步都有可核验的执行记录",
             ("典型用户", 20, NAVY, True),
             "外语专业学生、跨境电商/外贸从业者：频繁在中/英/俄/法/德之间切换",
             "手部不便或老年用户：需要按个人落点习惯调整键区",
             "创客与课堂：低成本国产开发板上的“AI + 硬件”完整样例"], size=17)

sl = slide("方案总览：KeyScreen 是什么")
arrow_flow(sl, ["触摸屏多点触控", "LubanCat 1H 渲染/判键/输入法", "USB HID 标准键盘", "电脑直接出字"], y=1.8)
arrow_flow(sl, ["用户一句话", "AGH + agnes-3.0-flash 规划", "MHS 设备工具调用", "键盘硬件被改造并复核"], y=3.3)
bullets(sl, ["下半屏：键盘（4 种语言、可调键位、可换壁纸配色）；上半屏：Agnes 对话框 + 猫娘桌宠 + 天气小组件",
             "电脑端：MHS 设备进程把键盘暴露为工具；面板负责 Agnes 对话桥接、Windows 键盘布局自动跟随、离线兜底",
             "所有 AI 操作都经过板端 REST API，有日志、可回放、可导出"], y=4.8, h=2.4, size=16)

sl = slide("系统架构")
rows = [("触摸屏 evdev", "多点触控坐标"), ("LubanCat 1H  board/keyscreen.py", "渲染 · 判键 · 输入法 · 桌宠 · REST API :8765"), ("USB HID gadget  /dev/hidg0", "标准键盘报文，任何系统免驱"), ("电脑", "记事本 / IDE 直接出字")]
for i, (a, b) in enumerate(rows):
    y = 1.7 + i * 1.3
    box(sl, Inches(0.6), Inches(y), Inches(5.6), Inches(1.0), NAVY if i % 2 == 0 else BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, Inches(0.8), Inches(y + 0.1), Inches(5.3), Inches(0.9), [(a, 15, WHITE, True), (b, 12, "D9E2F3")], gap=2)
    if i < 3:
        text(sl, Inches(3.1), Inches(y + 0.95), Inches(0.6), Inches(0.4), "↓", 20, GREY, True)
rows = [("Agnes agnes-3.0-flash", "意图理解 · 规划 · 工具调用 · 复核 · 回复"), ("Agnes Harness（AGH）", "会话 · 工具路由 · 执行记录导出（可脱敏）"), ("MHS 设备 keyscreen-01  board/mhs_device.py", "read_device / set_language / adjust_key / set_wallpaper / pet_act …"), ("desktop/panel.py", "对话桥接 · 布局跟随 · 会话轮换 · 离线兜底")]
for i, (a, b) in enumerate(rows):
    y = 1.7 + i * 1.3
    box(sl, Inches(7.0), Inches(y), Inches(5.8), Inches(1.0), "385723" if i % 2 == 0 else "548235", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, Inches(7.2), Inches(y + 0.1), Inches(5.5), Inches(0.9), [(a, 15, WHITE, True), (b, 12, "E2EFDA")], gap=2)
    if i < 3:
        text(sl, Inches(9.6), Inches(y + 0.95), Inches(0.6), Inches(0.4), "↕", 20, GREY, True)
text(sl, Inches(6.2), Inches(3.9), Inches(0.9), Inches(0.5), "⇄ 网络", 13, GREY, True, PP_ALIGN.CENTER)

sl = slide("硬件与 USB HID 键盘实现", "Linux configfs USB gadget：把开发板伪装成一把标准键盘")
cards(sl, [("硬件", ["LubanCat 1H（RK3566，国产开发板）", "电容触摸屏，evdev 多点触控", "USB-C OTG 口接电脑，网口接 PC 用于 AI 通道"]),
           ("HID 实现要点", ["configfs 创建 hid.usb0，标准 8 字节键盘报文", "强制 OTG 为 peripheral 模式，解决厂商守护进程抢占 UDC", "设备节点 major:minor 自检重建；写失败自动重开/重绑"]),
           ("效果", ["触摸到 HID 发送平均 1.75 ms，p95 3.3 ms", "支持组合键、长按连发、多指", "电脑复位 USB 后自动重连；开机 35 s 自启"])], size=14)

sl = slide("板端软件：渲染、判键与输入法")
cards(sl, [("渲染", ["pygame 脏矩形渲染，每帧只重绘变化区域", "定位桌面合成器瓶颈并关闭：刷新 48 ms → 27 ms", "壁纸、配色、键帽文字全部可由 API 热更新"]),
           ("判键", ["每键记录触点相对中心的 dx/dy，形成落点统计", "键位偏移可按键平移（钳制 ±40 px，防误操作）", "统计数据是 AI 校准与科学实验的原始数据"]),
           ("整句拼音输入法", ["4.7 万词条，词典由 jieba 词频 + pypinyin 生成", "Viterbi 整句转换，每键 < 1 ms；支持简拼 zgrm→中国人民", "用户词自学习；英文串原样保留"])], size=14)

sl = slide("创新点一：AI 真正“动手”改硬件", "通过 MHS 设备协议，把键盘的每个能力都变成 Agnes 可调用的工具")
table(sl, [("工具", "作用", "类型"),
           ("read_device", "读取语言、壁纸、配色、桌宠状态并附屏幕截图", "只读"),
           ("get_touch_stats", "读取每个键的按压次数、平均偏移、误触率", "只读"),
           ("adjust_key", "平移某个键的判定区域（±40 px 钳制）", "写硬件"),
           ("set_language", "切换 en/ru/fr/de 键盘，键帽即时换字", "写硬件"),
           ("set_wallpaper / set_profile / list_wallpapers", "更换壁纸与主题配色", "写外观"),
           ("pet_say / pet_act", "让桌宠说话、跳跃、挥手、欢呼", "交互")], y=1.9, size=13, widths=[4.2, 6.1, 2.0])
text(sl, Inches(0.5), Inches(5.3), Inches(12.3), Inches(1.5),
     ["每次调用都经过板端 REST API 校验（未知语言/壁纸/按键一律 400 拒绝），并写入日志；写操作有钳制与还原，AI 不可能把键盘改坏。"], 15, GREY)

sl = slide("创新点二：一句话多步定制（AGH 连续执行）", "用户：\“我要开始写俄语论文了，帮我把键盘配置好\" —— 真实运行 28–39 秒，6 次工具调用")
arrow_flow(sl, ["read_device\n读当前状态", "set_language\n切俄语", "list_wallpapers\n浏览壁纸", "set_wallpaper\nset_profile", "pet_act\ncheer", "read_device\n复核生效"], y=1.9, h=1.3, size=12)
bullets(sl, ["规划：Agnes 先读状态，再决定动哪些工具、用什么参数（选了 dark_ocean.png + 沉稳主题色）",
             "执行：连续调用设备工具，每一步的参数与返回值都在 AGH 执行记录中",
             "复核：再次读取设备状态确认 lang=ru、壁纸与配色已变化，然后用中俄双语回复用户",
             "电脑端同步：面板检测到语言变化后 1 秒内把 Windows 切到俄语布局，记事本直接打出俄文"], y=3.5, h=3.3, size=16)

sl = slide("AGH 执行记录：人工可核验", "evidence/submission/agh_ru_writing.html（AGH export，自动脱敏）")
pic(sl, "screen_after_ru_writing.jpg", 0.6, 1.7, w=6.2, cap="任务完成后的板端屏幕：俄语键帽 + 新壁纸")
bullets(sl, ["记录内容：用户输入、模型每一步的工具调用、参数、设备返回、最终回复",
             "配套：state_before.json / state_after.json 前后状态对比、屏幕截图、test_results.json",
             "一键复现：python3 tools/run_tests.py 重新跑一遍并重新导出",
             "工程证据目录另含 MHS 设备日志、AGH serve 日志、doctor 检查结果与多次历史运行"], x=7.1, y=1.8, w=5.7, h=5, size=15)

sl = slide("创新点三：四种语言键盘 + 电脑布局自动跟随")
pic(sl, "俄语键盘.jpg", 0.5, 1.6, w=6.0, cap="俄语 ЙЦУКЕН")
pic(sl, "法语键盘.jpg", 6.8, 1.6, w=6.0, cap="法语 AZERTY")
bullets(sl, ["英语 QWERTY · 俄语 ЙЦУКЕН · 法语 AZERTY · 德语 QWERTZ，键帽即时换字，所见即所得",
             "难点：HID 发送的是键位而非字符，电脑布局必须同步。面板轮询 /api/lang，用 LoadKeyboardLayout + WM_INPUTLANGCHANGEREQUEST 1 秒内切换，切回时恢复用户原输入法",
             "三条切换通道：屏幕快捷键（0.06 s 板端本地）、对 Agnes 说“换成法语键盘”、AGH 多步任务"], y=5.3, h=2.1, size=14)

sl = slide("创新点四：屏上 Agnes 助手与整句拼音输入法")
cards(sl, [("屏上就是对话框", ["键盘上半屏即 Agnes 聊天窗口，不占电脑屏幕", "快捷按钮：二次元模式 / 写代码配色 / 校准按键 / 换语言键盘 / 俄语写作 / 分析打字", "队列状态可见：排队、执行中、超时提示"]),
           ("搜狗式整句输入法", ["nihaoshijie → 你好世界；jintiantianqizenmeyang → 今天天气怎么样", "简拼 zgrm → 中国人民、wsm → 为什么", "keyscreen 等英文串原样保留，不强行转中文"]),
           ("为什么重要", ["Debian 10 + Python 3.7 的板子上没有现成中文输入法", "自研 Viterbi 整句转换每键 < 1 ms，不影响打字", "用户词自学习，越用越顺手"])], size=14)

sl = slide("创新点五：数据驱动的个性化键位校准")
pic(sl, "before.jpg", 0.5, 1.7, w=6.0, cap="校准前：Agnes 读取落点统计")
pic(sl, "after3.jpg", 6.8, 1.7, w=6.0, cap="校准后：偏差大的键被平移")
bullets(sl, ["每次按键记录相对键中心的偏移；Agnes 调用 get_touch_stats 看到“D 键平均偏 −11/+10 px”这类事实",
             "校准规则：取偏移一半、限 ±15 px、小步多次，可随时还原",
             "这是一个可被科学检验的假设——所以我们做了对照实验（下两页）"], y=5.3, h=2.1, size=14)

sl = slide("科学实验 ①：基线数据", "649 次板上真实按键（desktop/experiment.py baseline）")
stats(sl, [("649", "真实按键数"), ("22.1 px", "按键系统性偏差"), ("41%", "落在键边缘附近"), ("40 px", "平均偏移"), ("1.8%", "误触率")])
pic(sl, "baseline.png", 0.8, 3.6, h=3.3, cap="各键落点偏移分布")
bullets(sl, ["结论：偏差是系统性的（同一键总往同一方向偏），说明个性化校准在理论上有空间",
             "但偏差是否跨时间稳定？静态校准一次是否就够？需要留出法检验"], x=7.3, y=3.8, w=5.5, h=3, size=15)

sl = slide("科学实验 ②：留出法回放 —— 如实报告负结果", "按时间对半分：前 324 次学习每键偏移，后 325 次回放判定")
table(sl, [("指标", "校准前", "回放校准后"), ("误触率", "2.1%", "3.7%"), ("平均偏移", "34.0 px", "33.4 px"), ("按键系统偏差", "20.2 px", "18.8 px"), ("靠近键边缘", "38.5%", "37.5%")], x=0.6, y=1.9, w=5.6, size=14)
pic(sl, "replay.png", 6.6, 1.8, h=3.0)
bullets(sl, ["结果：静态一次性校准没有降低误触，偏移只下降 2%",
             "原因：跨会话落点漂移——坐姿、角度、心情都会改变落点，前半段学到的偏移到后半段已经不准",
             "结论与设计决策：校准必须由 Agnes 按会话持续、小步在线调整，而不是一次校完",
             "实验脚本、原始数据、图表全部随附，可完整复现；我们选择如实报告，而不是挑好看的数字"], y=4.9, h=2.4, size=14)

sl = slide("有温度的交互：桌宠、天气与壁纸")
pic(sl, "猫娘桌宠.jpg", 0.5, 1.6, w=6.0, cap="程序绘制的猫娘桌宠：可拖拽、会跑会跳、猫尾摆动、耳朵抖动")
pic(sl, "俄语写作模式.jpg", 6.8, 1.6, w=6.0, cap="俄语写作模式：Agnes 自选壁纸与沉稳配色")
bullets(sl, ["桌宠由代码绘制，Agnes 可通过 pet_act 让它 jump / wave / dance / cheer，成为 AI 操作结果的可见反馈",
             "苹果风格天气小组件（wttr.in），视频动态壁纸，“二次元模式 / 写代码配色”一键切换"], y=5.3, h=1.8, size=14)

sl = slide("可靠性工程：让演示现场不翻车")
cards(sl, [("板端", ["systemd 开机自启，冷启动 35 s 可用", "HID 写失败→重开设备→必要时重绑 UDC", "设备节点错乱自检重建；OTG 强制 peripheral"]),
           ("AI 通道", ["Agnes 会话每 30 分钟自动轮换，避免长会话卡死", "空回合/错误自动换新会话重试", "模型不可用 → 本地规则兜底回复（测试样例已验证）"]),
           ("输入与安全", ["所有 API 参数校验，未知值 400，偏移钳制 ±40 px", "密钥仅存本机 .agh/secrets/（600），仓库与提交包经扫描不含密钥", "语言切换三通道互为备份：本地快捷键 / Agnes / AGH"])], size=14)

sl = slide("实测指标（板上测量）")
table(sl, [("指标", "数值", "说明"),
           ("触摸 → USB HID 发送", "平均 1.75 ms，p95 3.3 ms", "板上时间戳测量"),
           ("按键 → 屏幕更新", "中位数约 34 ms", "脏矩形渲染"),
           ("屏幕刷新耗时", "48 ms → 27 ms", "关闭桌面合成器后"),
           ("拼音整句转换", "每键 < 1 ms", "4.7 万词条 Viterbi"),
           ("切换键盘语言", "板端 0.06 s，电脑布局 < 1 s", "三通道"),
           ("AGH 多步任务（6 次工具调用）", "28–39 s", "真实运行三次均完成"),
           ("Agnes 只读分析", "14–27 s", "get_touch_stats"),
           ("冷启动到可用", "约 35 s", "systemd 自启"),
           ("稳定性", "连续运行 1434 次按键，误触率 0.8%", "累计统计")], y=1.7, size=13, widths=[4.0, 4.3, 4.0])

sl = slide("测试样例：正常 / 边界 / 失败 共 11 条，全部通过", "真实设备运行，tools/run_tests.py 自动生成 运行证据/测试样例.md")
table(sl, [("类型", "用例", "结果"),
           ("正常", "AGH 多步任务“俄语写作”：语言=ru，壁纸/配色改变，复核通过", "通过 35.6 s"),
           ("正常", "快捷按钮换语言（板端本地通道）", "通过 0.06 s"),
           ("正常", "Agnes 只读分析打字数据（不改设置）", "通过 14 s"),
           ("正常", "整句拼音 nihaoshijie / zgrm / jintiantianqizenmeyang", "通过"),
           ("边界", "未知语言 xx → 400；空消息 → 400；偏移 999 → 钳制 40；英文串 keyscreen 原样", "4 条通过"),
           ("失败", "未知壁纸 → 400；未知按键 → 400", "通过"),
           ("失败", "模型不可用（--model no-such-model）→ AGH 退出码 2 → 本地离线兜底回复", "通过")], y=1.8, size=13, widths=[1.3, 8.6, 2.4])

sl = slide("我们解决了哪些问题")
table(sl, [("痛点", "KeyScreen 的做法", "证据"),
           ("外语键盘要买、要贴、切换易错", "一块屏四种语言，键帽即时换字，电脑布局 1 s 自动跟随并能恢复", "截图、测试样例、视频"),
           ("触屏键盘千人一面、误触", "逐键落点统计 + Agnes 在线小步校准；实验证明静态校准不够", "实验数据 baseline/replay"),
           ("AI 不会动手", "MHS 工具 + AGH 多步执行：读状态→改硬件→复核", "AGH 执行记录 html/jsonl"),
           ("板上没有中文输入法", "自研整句拼音输入法，每键 < 1 ms", "测试样例"),
           ("外设定制门槛高", "一句话定制：语言、壁纸、配色、桌宠一次配好", "俄语写作模式截图"),
           ("演示/长期运行不稳定", "自启、HID 重连、会话轮换、离线兜底", "冷启动与失败类测试")], y=1.7, size=13, widths=[3.4, 6.0, 2.9])

sl = slide("实际应用场景")
cards(sl, [("外语学习与考试", ["俄/法/德专业学生写作业、练打字，所见即所得", "教师机一键切换，无需给每台电脑贴键帽"]),
           ("跨境电商与外贸办公", ["客服在多语言之间秒切，电脑输入法自动跟随", "一句话切换“俄语客服模式”：语言 + 常用短语配色"]),
           ("无障碍与适老", ["按个人落点习惯调整键区、放大键帽", "AI 读取误触数据主动建议调整"]),
           ("创客教育", ["国产开发板上“AI + 硬件”完整开源样例", "REST API + MHS 工具，学生可扩展新技能"]),
           ("展厅与公共终端", ["多语言访客共用一块屏", "外观主题随活动一句话切换"]),
           ("开发者副屏", ["写代码配色 + 快捷键区", "Agnes 对话框常驻，不占主屏"])], h=2.3, cols=3, size=13)

sl = slide("社会价值与成本")
stats(sl, [("4 种", "语言键盘·一块屏"), ("≈ 300 元", "国产开发板 + 触摸屏"), ("0", "电脑端驱动安装"), ("100%", "AI 操作可审计")])
bullets(sl, ["降低多语言输入门槛：外语学习者、跨境从业者不再为每种语言买一把键盘",
             "数据驱动的无障碍：落点统计与 AI 调键位可服务手部不便、视力下降的用户",
             "国产硬件 + 开放协议：LubanCat（RK3566）+ AGH/MHS，可复制到其他外设（鼠标、手柄、控制面板）",
             "可信 AI 操作硬件的范式：工具白名单、参数钳制、执行记录导出、失败兜底"], y=3.7, h=3.3, size=16)

sl = slide("下一步计划")
cards(sl, [("近期（1 个月）", ["Agnes 按会话在线持续校准，并用 A/B 对照验证", "更多语言：日语假名、西班牙语、阿拉伯语", "语音唤起 Agnes，解放双手"]),
           ("中期（3 个月）", ["个人键盘画像：跨设备同步键位与偏好", "按应用自动切模式：IDE→代码配色，Word→写作模式", "手势与宏：滑动输入、自定义快捷键区"]),
           ("长期", ["开放 MHS 技能市场，让社区为键盘写新工具", "无障碍模式认证与用户研究", "量产级外壳与一体化模组"])], size=14)

sl = slide("团队与声明")
cards(sl, [("贾灵鑫（队长）", ["方案与架构设计", "板端开发：USB HID、判键、多语言", "AGH/MHS 接入与多步任务", "实验设计与数据分析"]),
           ("队员二【姓名】", ["电脑端：MHS 设备、对话桥接、布局跟随", "输入法与屏上助手界面", "桌宠、壁纸与视觉设计", "视频、公开内容与答辩准备"]),
           ("诚信声明", ["第三方素材与数据来源已逐项声明", "开发中使用 Cursor AI 编程助手辅助，全部代码经队员审阅并上板验证", "运行证据、测试与实验数据均来自真实设备，未经伪造", "比赛后轮换 API 密钥"])], size=14)

sl = prs.slides.add_slide(BL)
n[0] += 1
box(sl, 0, 0, W, H, NAVY)
text(sl, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2), "让键盘随人而变", 48, WHITE, True, PP_ALIGN.CENTER)
text(sl, Inches(0.8), Inches(3.5), Inches(11.5), Inches(1), "KeyScreen · 无敌键队 · U376", 24, "BDD7EE", align=PP_ALIGN.CENTER)
text(sl, Inches(0.8), Inches(4.6), Inches(11.5), Inches(1.4), ["谢谢评委老师，欢迎上手体验", "现场可演示：插线打字 → 一句话俄语写作模式 → AGH 执行记录"], 18, WHITE, align=PP_ALIGN.CENTER, gap=8)

prs.save(sys.argv[1])
print("slides", n[0])
