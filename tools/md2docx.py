"""Minimal Markdown -> .docx for the submission docs. Usage: md2docx.py IN.md|IN.txt OUT.docx"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK = "微软雅黑"


def font(run, size=None, bold=None, color=None, mono=False):
    run.font.name = "Consolas" if mono else CJK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def inline(par, text, size=11):
    """**bold**, `code`, 【placeholder】 highlighted."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|【[^】]*】|_{4,})", text):
        if not tok:
            continue
        if tok.startswith("**"):
            font(par.add_run(tok[2:-2]), size, bold=True)
        elif tok.startswith("`"):
            font(par.add_run(tok[1:-1]), size - 1, mono=True, color="1F4E79")
        elif tok.startswith("【"):
            font(par.add_run(tok), size, bold=True, color="C00000")
        else:
            font(par.add_run(tok), size)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\s*\|?\s*:?-{2,}", r)]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            p = t.cell(i, j).paragraphs[0]
            inline(p, c, 9.5)
            if i == 0:
                for run in p.runs:
                    run.font.bold = True
    doc.add_paragraph()


def convert(src, dst):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = CJK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    st.font.size = Pt(11)
    lines = open(src, encoding="utf-8").read().splitlines()
    i, code, buf = 0, False, []
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            if code:
                p = doc.add_paragraph()
                font(p.add_run("\n".join(buf)), 9.5, mono=True)
                buf = []
            code = not code
            i += 1
            continue
        if code:
            buf.append(ln)
            i += 1
            continue
        if ln.lstrip().startswith("|"):
            j = i
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                j += 1
            table(doc, lines[i:j])
            i = j
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            lvl = len(m.group(1))
            p = doc.add_paragraph()
            font(p.add_run(m.group(2)), [20, 15, 13][lvl - 1], bold=True, color="1F3864")
            if lvl == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r"^【.+】$", ln.strip()):  # 作品详情.txt section headers
            p = doc.add_paragraph()
            font(p.add_run(ln.strip()), 14, bold=True, color="1F3864")
        elif re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet")
            inline(p, re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+[.、]\s*", ln):
            p = doc.add_paragraph(style="List Number")
            inline(p, re.sub(r"^\s*\d+[.、]\s*", "", ln))
        elif ln.startswith(">"):
            p = doc.add_paragraph()
            inline(p, ln.lstrip("> "), 10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif ln.strip():
            p = doc.add_paragraph()
            inline(p, ln)
            if ln.startswith("签名"):
                p.paragraph_format.space_before = Pt(24)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
